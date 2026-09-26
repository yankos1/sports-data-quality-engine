import json
import logging
import os
import tempfile
from pathlib import Path
from typing import Optional

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import text

from src.analysis.drift_analyzer import OddsDriftAnalyzer
from src.database.connection import SessionLocal, engine
from src.reporting.clv_analyzer import CLVAnalyzer
from src.reporting.drift_analyzer import MarketDriftAnalyzer


LOGGER = logging.getLogger(__name__)
PROJECT_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_OUTPUT_PATH = PROJECT_ROOT / "data" / "reports" / "dq_report_latest.xlsx"

CLEAN_ODDS_QUERY = text(
    """
    SELECT
        f.kickoff_at AS `Date coup d'envoi UTC`,
        home_team.canonical_name AS `Équipe Domicile`,
        away_team.canonical_name AS `Équipe Extérieur`,
        clean_odds.bookmaker AS Bookmaker,
        clean_odds.odds_home AS `Cote 1`,
        clean_odds.odds_draw AS `Cote N`,
        clean_odds.odds_away AS `Cote 2`,
        (1.0 / clean_odds.odds_home
         + 1.0 / clean_odds.odds_draw
         + 1.0 / clean_odds.odds_away) * 100 AS `Overround (%)`,
        clean_odds.captured_at AS `Horodatage d'ingestion`
    FROM clean_odds
    JOIN (
        SELECT fixture_id, bookmaker, MAX(captured_at) AS captured_at
        FROM clean_odds
        GROUP BY fixture_id, bookmaker
    ) AS latest_odds
        ON latest_odds.fixture_id = clean_odds.fixture_id
        AND latest_odds.bookmaker = clean_odds.bookmaker
        AND latest_odds.captured_at = clean_odds.captured_at
    JOIN fixtures AS f ON f.id = clean_odds.fixture_id
    JOIN teams AS home_team ON home_team.id = f.home_team_id
    JOIN teams AS away_team ON away_team.id = f.away_team_id
    ORDER BY f.kickoff_at, home_team.canonical_name
    """
)

DQ_EXCEPTIONS_QUERY = text(
    """
    SELECT
        id AS ID,
        created_at AS Horodatage,
        rule_code AS `Code Règle`,
        message AS Message,
        source AS `Source/Bookmaker`,
        rejected_payload AS `Payload JSON brut`
    FROM dq_exceptions
    ORDER BY created_at DESC, id DESC
    """
)


class ExcelReportGenerator:
    def generate_report(self, output_path: Optional[Path] = None) -> Path:
        report_path = Path(output_path) if output_path is not None else DEFAULT_OUTPUT_PATH
        report_path = report_path.expanduser().resolve()
        report_path.parent.mkdir(parents=True, exist_ok=True)

        clean_odds_df, exceptions_df, drift_df, clv_df, odds_drift_df = (
            self._load_dataframes()
        )
        odds_drift_df, clv_beat_rate = self._prepare_odds_drift(odds_drift_df)
        total_clean_odds = len(clean_odds_df.index)
        total_exceptions = len(exceptions_df.index)
        total_analyzed = total_clean_odds + total_exceptions
        quality_score = (
            total_clean_odds / total_analyzed if total_analyzed else 0.0
        )
        rule_counts = (
            exceptions_df["Code Règle"].fillna("UNKNOWN").value_counts().sort_index()
            if total_exceptions
            else pd.Series(dtype="int64")
        )
        summary_metrics = pd.DataFrame(
            [
                ("Total cotes analysées", total_analyzed),
                ("Cotes conformes (clean_odds)", total_clean_odds),
                ("Anomalies (dq_exceptions)", total_exceptions),
                ("Data Quality Score", quality_score),
                ("Taux de battement CLV (%)", clv_beat_rate),
            ],
            columns=["Métrique", "Valeur"],
        )
        anomaly_counts = pd.DataFrame(
            [(rule_code, int(count)) for rule_code, count in rule_counts.items()],
            columns=["Code Règle", "Nombre d'anomalies"],
        )

        exceptions_df = exceptions_df.copy()
        if "Payload JSON brut" in exceptions_df.columns:
            exceptions_df["Payload JSON brut"] = exceptions_df[
                "Payload JSON brut"
            ].map(self._format_payload)
        if "Variation (%)" in drift_df.columns:
            drift_df = drift_df.copy()
            drift_df["Variation (%)"] = drift_df["Variation (%)"] / 100
        if not clv_df.empty:
            clv_df = clv_df.copy()
            clv_df["CLV (%)"] = clv_df["CLV (%)"] / 100
            clv_df["Gain Fair Proba (pts)"] = clv_df["Gain Fair Proba (pts)"] / 100
        if not odds_drift_df.empty:
            odds_drift_df = odds_drift_df.copy()
            odds_drift_df["Drift 1 (%)"] = odds_drift_df["Drift 1 (%)"] / 100
            odds_drift_df["Drift 2 (%)"] = odds_drift_df["Drift 2 (%)"] / 100

        file_descriptor, temporary_name = tempfile.mkstemp(
            prefix="{}-".format(report_path.stem),
            suffix=report_path.suffix,
            dir=report_path.parent,
        )
        os.close(file_descriptor)
        temporary_path = Path(temporary_name)
        try:
            with pd.ExcelWriter(temporary_path, engine="openpyxl") as writer:
                summary_metrics.to_excel(
                    writer,
                    sheet_name="Synthèse DQ",
                    startrow=2,
                    index=False,
                )
                anomaly_counts.to_excel(
                    writer,
                    sheet_name="Synthèse DQ",
                    startrow=9,
                    index=False,
                )
                clean_odds_df.to_excel(writer, sheet_name="Clean Odds", index=False)
                exceptions_df.to_excel(writer, sheet_name="DQ Exceptions Log", index=False)
                drift_df.to_excel(writer, sheet_name="Market Trends & Drift", index=False)
                clv_df.to_excel(writer, sheet_name="CLV Backtest Analysis", index=False)
                odds_drift_df.to_excel(writer, sheet_name="Analyse Odds Drift", index=False)

                summary_sheet = writer.book["Synthèse DQ"]
                summary_sheet.merge_cells("A1:B1")
                summary_sheet["A1"] = "Synthèse DQ"
                summary_sheet["A9"] = "Anomalies par règle"
                self._style_workbook(writer.book)

            os.replace(temporary_path, report_path)
        except PermissionError as error:
            LOGGER.exception("Le rapport Excel est verrouillé ou inaccessible: %s", report_path)
            raise PermissionError(
                "Impossible de remplacer le rapport; fermez le classeur s'il est ouvert: {}".format(
                    report_path
                )
            ) from error
        finally:
            if temporary_path.exists():
                try:
                    temporary_path.unlink()
                except OSError:
                    LOGGER.warning("Impossible de supprimer le fichier temporaire %s", temporary_path)

        LOGGER.info("Rapport qualité généré : %s", report_path)
        return report_path

    @staticmethod
    def _load_dataframes() -> tuple[
        pd.DataFrame,
        pd.DataFrame,
        pd.DataFrame,
        pd.DataFrame,
        pd.DataFrame,
    ]:
        with engine.connect() as connection:
            clean_odds_df = pd.read_sql(CLEAN_ODDS_QUERY, connection)
            exceptions_df = pd.read_sql(DQ_EXCEPTIONS_QUERY, connection)
        drift_df = MarketDriftAnalyzer().get_significant_drifts()
        clv_df = CLVAnalyzer().get_historical_clv()
        with SessionLocal() as session:
            odds_analyzer = OddsDriftAnalyzer(session)
            opening_closing = odds_analyzer.get_opening_and_closing_odds()
            odds_drift_df = odds_analyzer.calculate_clv_metrics(opening_closing)
        return clean_odds_df, exceptions_df, drift_df, clv_df, odds_drift_df

    @staticmethod
    def _prepare_odds_drift(
        metrics: pd.DataFrame,
    ) -> tuple[pd.DataFrame, object]:
        output_columns = [
            "Date coup d'envoi UTC",
            "Compétition",
            "Équipe Domicile",
            "Équipe Extérieur",
            "Cote 1 (Ouverture)",
            "Cote 1 (Actuelle)",
            "Drift 1 (%)",
            "Cote 2 (Ouverture)",
            "Cote 2 (Actuelle)",
            "Drift 2 (%)",
            "CLV Favorable (Oui/Non)",
            "Nb Snapshots capturés",
        ]
        if metrics.empty:
            return pd.DataFrame(columns=output_columns), "Aucun snapshot disponible"

        favorite_columns = ["open_odds_home", "open_odds_draw", "open_odds_away"]
        outcomes = {"open_odds_home": "home", "open_odds_draw": "draw", "open_odds_away": "away"}
        favorite_keys = metrics[favorite_columns].idxmin(axis=1)
        open_favorite = [
            row[key] for (_, row), key in zip(metrics.iterrows(), favorite_keys)
        ]
        close_favorite = [
            row["close_odds_{}".format(outcomes[key])]
            for (_, row), key in zip(metrics.iterrows(), favorite_keys)
        ]
        snapshot_counts = pd.to_numeric(metrics["snapshot_count"], errors="coerce").fillna(0)
        clv_won = [opening > closing for opening, closing in zip(open_favorite, close_favorite)]
        statuses = [
            "En attente de snapshots (T0 unique)"
            if snapshot_count < 2
            else ("Oui" if beaten else "Non")
            for snapshot_count, beaten in zip(snapshot_counts, clv_won)
        ]
        evaluable = snapshot_counts >= 2
        beat_rate = (
            float(pd.Series(clv_won, index=metrics.index).loc[evaluable].mean())
            if evaluable.any()
            else "En attente de snapshots (T0 unique)"
        )

        odds_drift_df = pd.DataFrame(
            {
                "Date coup d'envoi UTC": metrics.get("kickoff_at"),
                "Compétition": metrics.get("competition"),
                "Équipe Domicile": metrics.get("home_team"),
                "Équipe Extérieur": metrics.get("away_team"),
                "Cote 1 (Ouverture)": metrics["open_odds_home"],
                "Cote 1 (Actuelle)": metrics["close_odds_home"],
                "Drift 1 (%)": metrics["drift_home_pct"],
                "Cote 2 (Ouverture)": metrics["open_odds_away"],
                "Cote 2 (Actuelle)": metrics["close_odds_away"],
                "Drift 2 (%)": metrics["drift_away_pct"],
                "CLV Favorable (Oui/Non)": statuses,
                "Nb Snapshots capturés": snapshot_counts.astype(int),
            },
            index=metrics.index,
        )
        return odds_drift_df[output_columns], beat_rate

    @staticmethod
    def _format_payload(payload: object) -> str:
        if payload is None:
            return ""
        if isinstance(payload, str):
            try:
                payload = json.loads(payload)
            except json.JSONDecodeError:
                return payload
        return json.dumps(payload, ensure_ascii=False, default=str)

    @classmethod
    def _style_workbook(cls, workbook: Workbook) -> None:
        header_fill = PatternFill(fill_type="solid", fgColor="243447")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        title_fill = PatternFill(fill_type="solid", fgColor="15616D")
        title_font = Font(name="Calibri", size=15, bold=True, color="FFFFFF")
        thin_gray = Side(style="thin", color="D5DCE3")
        cell_border = Border(
            left=thin_gray,
            right=thin_gray,
            top=thin_gray,
            bottom=thin_gray,
        )

        for worksheet in workbook.worksheets:
            worksheet.freeze_panes = "A2" if worksheet.title != "Synthèse DQ" else "A4"
            worksheet.sheet_view.showGridLines = False

            for row in worksheet.iter_rows():
                for cell in row:
                    cell.border = cell_border
                    cell.alignment = Alignment(vertical="top", wrap_text=True)

            header_rows = [1]
            if worksheet.title == "Synthèse DQ":
                header_rows = [3, 10]
                worksheet["A1"].fill = title_fill
                worksheet["A1"].font = title_font
                worksheet["A1"].alignment = Alignment(vertical="center")
                worksheet.row_dimensions[1].height = 26

            for header_row in header_rows:
                if header_row > worksheet.max_row:
                    continue
                for cell in worksheet[header_row]:
                    if cell.value is not None:
                        cell.fill = header_fill
                        cell.font = header_font
                        cell.alignment = Alignment(
                            horizontal="left",
                            vertical="center",
                            wrap_text=True,
                        )

            if worksheet.title == "Synthèse DQ":
                for row_index in range(4, worksheet.max_row + 1):
                    if worksheet.cell(row_index, 1).value in (
                        "Data Quality Score",
                        "Taux de battement CLV (%)",
                    ):
                        worksheet.cell(row_index, 2).number_format = "0.00%"

            if worksheet.title == "Clean Odds":
                for row_index in range(2, worksheet.max_row + 1):
                    worksheet.cell(row_index, 8).number_format = '0.00"%"'

            if worksheet.title == "Market Trends & Drift":
                header_columns = {
                    cell.value: cell.column for cell in worksheet[1] if cell.value
                }
                variation_column = header_columns["Variation (%)"]
                column_letter = get_column_letter(variation_column)
                green_fill = PatternFill(fill_type="solid", fgColor="D9EAD3")
                orange_fill = PatternFill(fill_type="solid", fgColor="FCE4D6")
                reversal_fill = PatternFill(fill_type="solid", fgColor="E4DFEC")
                reversal_font = Font(bold=True, color="7030A0")
                for row_index in range(2, worksheet.max_row + 1):
                    worksheet.cell(row_index, variation_column).number_format = (
                        "+0.00%;-0.00%"
                    )
                    for probability_column in (
                        "Ancienne proba (%)",
                        "Nouvelle proba (%)",
                    ):
                        if probability_column in header_columns:
                            worksheet.cell(
                                row_index,
                                header_columns[probability_column],
                            ).number_format = "0.00%"
                    if "Gain de proba (pts)" in header_columns:
                        worksheet.cell(
                            row_index,
                            header_columns["Gain de proba (pts)"],
                        ).number_format = "+0.00%;-0.00%"
                    if "Cote Fair actuelle" in header_columns:
                        worksheet.cell(
                            row_index,
                            header_columns["Cote Fair actuelle"],
                        ).number_format = "0.00"
                    if "Reversal Intensity (%)" in header_columns:
                        worksheet.cell(
                            row_index,
                            header_columns["Reversal Intensity (%)"],
                        ).number_format = "0.00\"%\""
                if worksheet.max_row > 1:
                    drift_range = "{}2:{}{}".format(
                        column_letter,
                        column_letter,
                        worksheet.max_row,
                    )
                    worksheet.conditional_formatting.add(
                        drift_range,
                        FormulaRule(
                            formula=["${}2<0".format(column_letter)],
                            fill=green_fill,
                        ),
                    )
                    if "Signal Marché" in header_columns:
                        signal_letter = get_column_letter(header_columns["Signal Marché"])
                        signal_range = "{}2:{}{}".format(
                            signal_letter,
                            signal_letter,
                            worksheet.max_row,
                        )
                        worksheet.conditional_formatting.add(
                            signal_range,
                            FormulaRule(
                                formula=['LEFT(${0}2,14)="SHARP REVERSAL"'.format(signal_letter)],
                                fill=reversal_fill,
                                font=reversal_font,
                            ),
                        )
                    worksheet.conditional_formatting.add(
                        drift_range,
                        FormulaRule(
                            formula=["${}2>0".format(column_letter)],
                            fill=orange_fill,
                        ),
                    )

            if worksheet.title == "CLV Backtest Analysis":
                header_columns = {
                    cell.value: cell.column for cell in worksheet[1] if cell.value
                }
                if "CLV (%)" in header_columns:
                    clv_column = header_columns["CLV (%)"]
                    for row_index in range(2, worksheet.max_row + 1):
                        worksheet.cell(row_index, clv_column).number_format = (
                            "+0.00%;-0.00%"
                        )
                        worksheet.cell(
                            row_index,
                            header_columns["Gain Fair Proba (pts)"],
                        ).number_format = "+0.00%;-0.00%"
                        for odds_column in ("Cote Prise", "Cote Clôture"):
                            worksheet.cell(
                                row_index,
                                header_columns[odds_column],
                            ).number_format = "0.00"

                if "CLV (%)" in header_columns and worksheet.max_row > 1:
                    clv_letter = get_column_letter(clv_column)
                    clv_range = "{}2:{}{}".format(
                        clv_letter,
                        clv_letter,
                        worksheet.max_row,
                    )
                    worksheet.conditional_formatting.add(
                        clv_range,
                        FormulaRule(
                            formula=["${}2>0".format(clv_letter)],
                            fill=PatternFill(fill_type="solid", fgColor="D9EAD3"),
                        ),
                    )
                    worksheet.conditional_formatting.add(
                        clv_range,
                        FormulaRule(
                            formula=["${}2<0".format(clv_letter)],
                            fill=PatternFill(fill_type="solid", fgColor="FCE8E6"),
                        ),
                    )

            if worksheet.title == "Analyse Odds Drift":
                header_columns = {
                    cell.value: cell.column for cell in worksheet[1] if cell.value
                }
                for row_index in range(2, worksheet.max_row + 1):
                    for drift_column in ("Drift 1 (%)", "Drift 2 (%)"):
                        worksheet.cell(
                            row_index,
                            header_columns[drift_column],
                        ).number_format = "+0.00%;-0.00%"
                    for odds_column in (
                        "Cote 1 (Ouverture)",
                        "Cote 1 (Actuelle)",
                        "Cote 2 (Ouverture)",
                        "Cote 2 (Actuelle)",
                    ):
                        worksheet.cell(
                            row_index,
                            header_columns[odds_column],
                        ).number_format = "0.00"
                    worksheet.cell(
                        row_index,
                        header_columns["Nb Snapshots capturés"],
                    ).number_format = "0"

                if worksheet.max_row > 1:
                    clv_column = get_column_letter(
                        header_columns["CLV Favorable (Oui/Non)"]
                    )
                    clv_range = "{}2:{}{}".format(
                        clv_column,
                        clv_column,
                        worksheet.max_row,
                    )
                    worksheet.conditional_formatting.add(
                        clv_range,
                        FormulaRule(
                            formula=['${}2="Oui"'.format(clv_column)],
                            fill=PatternFill(fill_type="solid", fgColor="D9EAD3"),
                        ),
                    )
                    worksheet.conditional_formatting.add(
                        clv_range,
                        FormulaRule(
                            formula=['${}2="Non"'.format(clv_column)],
                            fill=PatternFill(fill_type="solid", fgColor="FCE8E6"),
                        ),
                    )

            cls._adjust_column_widths(worksheet)

        exception_sheet = workbook["DQ Exceptions Log"]
        if exception_sheet.max_row > 1:
            critical_fill = PatternFill(fill_type="solid", fgColor="FCE8E6")
            exception_sheet.conditional_formatting.add(
                "A2:F{}".format(exception_sheet.max_row),
                FormulaRule(
                    formula=['LEFT($C2,7)="DQ_ERR_"'],
                    fill=critical_fill,
                ),
            )

    @staticmethod
    def _adjust_column_widths(worksheet) -> None:
        for column_cells in worksheet.columns:
            column_letter = get_column_letter(column_cells[0].column)
            max_length = max(
                (len(str(cell.value)) for cell in column_cells if cell.value is not None),
                default=0,
            )
            worksheet.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 70)


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")
    generated_path = ExcelReportGenerator().generate_report()
    print(generated_path)